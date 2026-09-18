#!/usr/bin/env python3
"""Build group-meeting PPT for recent A/B pipeline + latency analysis."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import nsmap, qn
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).with_name("组会汇报_感知决策动作生成_202608.pdf")
# pptx path
OUT_PPTX = Path(__file__).with_name("组会汇报_感知决策动作生成_202608.pptx")

NAVY = RGBColor(0x0B, 0x2A, 0x4A)
NAVY2 = RGBColor(0x12, 0x3A, 0x63)
ACCENT = RGBColor(0x1F, 0x7A, 0x8C)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF4, 0xF7, 0xFA)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x4A, 0x55, 0x63)
RED = RGBColor(0xB4, 0x23, 0x18)
GREEN = RGBColor(0x1B, 0x7A, 0x3A)
ORANGE = RGBColor(0xC4, 0x5C, 0x12)

FONT = "PingFang SC"
W, H = Inches(13.333), Inches(7.5)


def set_run(run, text, size=16, bold=False, color=INK, font=FONT):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", font)


def add_rect(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def add_text(slide, l, t, w, h, text, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    set_run(run, text, size, bold, color)
    return box


def _p(tf, text, size=15, bold=False, color=INK, space=6):
    if not tf.paragraphs[0].runs:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.space_after = Pt(space)
    p.level = 0
    run = p.add_run()
    set_run(run, text, size, bold, color)
    return p


def bullets(slide, l, t, w, h, items, size=16, color=INK):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        p.level = 0
        if isinstance(item, tuple):
            txt, b = item
            run = p.add_run()
            set_run(run, txt, size, b, color)
        else:
            run = p.add_run()
            set_run(run, item, size, False, color)
    return box


def header(slide, title, subtitle=""):
    add_rect(slide, 0, 0, W, Inches(1.05), NAVY)
    add_rect(slide, 0, Inches(1.05), W, Inches(0.06), GOLD)
    add_text(slide, Inches(0.45), Inches(0.18), Inches(12.4), Inches(0.5), title, 26, True, WHITE)
    if subtitle:
        add_text(slide, Inches(0.45), Inches(0.62), Inches(12.4), Inches(0.35), subtitle, 13, False, RGBColor(0xC5, 0xD4, 0xE0))
    add_rect(slide, 0, Inches(7.28), W, Inches(0.22), NAVY2)
    add_text(
        slide,
        Inches(0.45),
        Inches(7.28),
        Inches(12.4),
        Inches(0.22),
        "HKUST RBM  ·  社交机器人感知→决策→动作生成  ·  组会 2026.08",
        10,
        False,
        WHITE,
    )


def table_slide(slide, rows, left, top, width, row_h=0.38, col_w=None, header_fill=NAVY, font_size=13):
    n_rows, n_cols = len(rows), len(rows[0])
    if col_w is None:
        col_w = [width / n_cols] * n_cols
    tbl_shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, Inches(row_h * n_rows))
    tbl = tbl_shape.table
    for i, w in enumerate(col_w):
        tbl.columns[i].width = int(w)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
            run = p.add_run()
            set_run(run, str(val), font_size, r == 0, WHITE if r == 0 else INK)
            fill = header_fill if r == 0 else (LIGHT if r % 2 else WHITE)
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
    return tbl


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, W, H, WHITE)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # 0 title
    s = new_slide(prs)
    add_rect(s, 0, 0, W, H, NAVY)
    add_rect(s, 0, Inches(5.9), W, Inches(1.6), NAVY2)
    add_rect(s, Inches(0.55), Inches(2.05), Inches(1.4), Inches(0.08), GOLD)
    add_text(s, Inches(0.55), Inches(1.15), Inches(12), Inches(0.45), "组会进展汇报", 18, False, GOLD)
    add_text(
        s,
        Inches(0.55),
        Inches(2.25),
        Inches(12.2),
        Inches(1.4),
        "感知 → 决策 → MoMask 动作生成",
        36,
        True,
        WHITE,
    )
    add_text(
        s,
        Inches(0.55),
        Inches(3.7),
        Inches(12.2),
        Inches(1.0),
        "A∥B 双轨架构落地 · 云侧联调计时拆解 · 端侧 MoMask 实测 · Qwen 部署进展",
        18,
        False,
        RGBColor(0xC5, 0xD4, 0xE0),
    )
    add_text(s, Inches(0.55), Inches(6.15), Inches(12), Inches(0.4), "HKUST  ·  RBM Project  ·  2026.08", 16, False, WHITE)
    add_text(s, Inches(0.55), Inches(6.55), Inches(12), Inches(0.35), "覆盖周期：基线冻结、Mac 云侧 E2E、树莓派 4GB 部署与内存约束", 13, False, RGBColor(0xA8, 0xBC, 0xCC))

    # 1 agenda
    s = new_slide(prs)
    header(s, "汇报提纲", "本次组会按「架构 → 实验 → 延迟拆解 → 端侧」四段讲")
    items = [
        ("01  目标链路与 A∥B 架构（已冻结）", "双轨 + 薄仲裁，默认 A 并行 B"),
        ("02  本阶段完成项与基线回归", "预置动作 enrichment、Mac pipeline、冻结标签"),
        ("03  云侧联调实验结果与延迟拆解", "为何总时延 15s；视觉帧数 / 听觉处理"),
        ("04  MoMask 2s vs 8.9s：口径问题，不是模型退化", "Pi 热启动 vs Mac 冷启动 subprocess"),
        ("05  端侧部署与下一步", "Qwen2.5-1.5B、4GB 编译 OOM、16GB 真机策略"),
    ]
    for i, (t, d) in enumerate(items):
        y = Inches(1.35 + i * 1.05)
        add_rect(s, Inches(0.5), y, Inches(12.3), Inches(0.92), LIGHT)
        add_rect(s, Inches(0.5), y, Inches(0.12), Inches(0.92), ACCENT)
        add_text(s, Inches(0.85), y + Inches(0.12), Inches(11.7), Inches(0.4), t, 20, True, NAVY)
        add_text(s, Inches(0.85), y + Inches(0.48), Inches(11.7), Inches(0.35), d, 14, False, MUTED)

    # 2 goal
    s = new_slide(prs)
    header(s, "项目目标链路", "职责边界：本侧做感知→决策→关节；队友做真机执行映射")
    add_text(s, Inches(0.5), Inches(1.3), Inches(12.3), Inches(0.4), "正式闭环（从后往前对齐）", 16, True, NAVY)
    stages = [
        ("感知", "FaceMesh 几何情绪\n+ 麦克风 ASR"),
        ("Decide", "结构化 JSON\n→ 英文 action_prompt"),
        ("Arbiter", "A_only / B_only\n/ A_parallel_B"),
        ("生成", "A: ActionGroup\nB: MoMask joints"),
        ("执行", "队友重定向\n舵机 / 关节"),
    ]
    for i, (t, d) in enumerate(stages):
        x = Inches(0.45 + i * 2.55)
        add_rect(s, x, Inches(1.9), Inches(2.35), Inches(2.15), NAVY if i != 2 else ACCENT)
        add_text(s, x, Inches(2.05), Inches(2.35), Inches(0.5), t, 18, True, WHITE, PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.1), Inches(2.6), Inches(2.15), Inches(1.2), d, 13, False, WHITE, PP_ALIGN.CENTER)
        if i < 4:
            add_text(s, x + Inches(2.15), Inches(2.6), Inches(0.4), Inches(0.4), "→", 22, True, GOLD, PP_ALIGN.CENTER)
    bullets(
        s,
        Inches(0.5),
        Inches(4.3),
        Inches(12.3),
        Inches(2.6),
        [
            "测试机：树莓派 5 · 4GB（联调 / 压测）　　真机：树莓派 5 · 16GB（机器人内）",
            "MoMask 吃的是英文 HumanML3D 句式，不是 4 类情绪标签本身",
            "原则：A 保证「立刻能动、可安全停下」；B 保证「表达细、出 (T,22,3) 关节」",
            "失败落点永远在 A：B 挂了不撤销已执行的预置动作",
        ],
        16,
    )

    # 3 architecture
    s = new_slide(prs)
    header(s, "A∥B 架构（已冻结，2026-08-18）", "标签 baseline/mac-a-preset-momask-v1　·　文档 BASELINE_FREEZE.md")
    rows = [
        ["轨道", "角色", "输入", "输出", "时延目标"],
        ["A 预置", "保底 / 快路径", "粗情绪 + 强度", "ActionGroup 名", "启动 < 0.5s"],
        ["B 正式", "表达主路径", "action_prompt", "(T,22,3) joints.npy", "热启动 ~2.3s（Pi）"],
        ["Arbiter", "薄仲裁，不内嵌模型", "mode + conf + 错误码", "本轮走 A / B / 并行", "可忽略"],
    ]
    table_slide(
        s,
        rows,
        Inches(0.45),
        Inches(1.3),
        Inches(12.4),
        0.42,
        [Inches(1.8), Inches(2.4), Inches(2.6), Inches(2.8), Inches(2.8)],
        font_size=13,
    )
    bullets(
        s,
        Inches(0.5),
        Inches(3.55),
        Inches(12.3),
        Inches(3.4),
        [
            ("默认模式：A_parallel_B —— 先播预置，同时/随后跑 MoMask。", True),
            "降级规则：confidence < 0.45 或 fallback 或空 prompt 或 Decide/MoMask 失败 → 本轮只 A。",
            "三模式可配置：A_only（演示保底） / B_only（给队友交关节包） / A_parallel_B（真机默认）。",
            "精细 prompt 由 Decide 写；4 类情绪只服务 A 选组和仲裁，不能单独当 MoMask 输入。",
        ],
        16,
    )

    # 4 this period
    s = new_slide(prs)
    header(s, "本阶段完成项（对照缺口）", "上周期：MoMask 端侧可跑；本周期：把「前面」感知与决策接上")
    cards = [
        ("轨道 A", GREEN, ["强度分档 mild/strong（阈值 0.55）", "扩映射 + 禁攻击动作", "离线测试 ALL PASS", "Mac 策略监控 Demo"]),
        ("轨道 B / 联调", ACCENT, ["pipeline/ Arbiter + Track A/B", "Mac 计时 E2E + 实时窗口", "云 Decide 通路打通", "决策 mock 冒烟 20/20"]),
        ("端侧", GOLD, ["Pi 100 条 ~2.30s/条", "代码已 rsync 到 4GB Pi", "Decide 改为 Qwen 路线", "规则版 A∥B dry-run 已通"]),
        ("约束与坑", RED, ["云 ASR whisper-1 不可用", "每轮开摄像头标定 ~10s", "MoMask subprocess 冷启动", "4GB 并行编译会 OOM"]),
    ]
    for i, (title, color, lines) in enumerate(cards):
        x = Inches(0.4 + (i % 4) * 3.2)
        y = Inches(1.35)
        add_rect(s, x, y, Inches(3.05), Inches(5.55), LIGHT)
        add_rect(s, x, y, Inches(3.05), Inches(0.55), color)
        add_text(s, x, y + Inches(0.1), Inches(3.05), Inches(0.4), title, 18, True, WHITE, PP_ALIGN.CENTER)
        box = s.shapes.add_textbox(x + Inches(0.15), y + Inches(0.75), Inches(2.75), Inches(4.5))
        tf = box.text_frame
        tf.word_wrap = True
        for j, line in enumerate(lines):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.space_after = Pt(12)
            run = p.add_run()
            set_run(run, "•  " + line, 14, False, INK)

    # 5 track A mapping
    s = new_slide(prs)
    header(s, "轨道 A：预置动作 enrichment", "权威文件 EmotionActionScheduler.py　·　回归 preset_action_tests")
    rows = [
        ["情绪", "mild（conf < 0.55）", "strong（conf ≥ 0.55）"],
        ["neutral", "无动作", "无动作"],
        ["happy", "wave / stepping", "chest / wave / twist"],
        ["unhappy", "bow / jugong", "squat / bow / jugong"],
        ["surprised", "twist / stepping", "back_fast / twist / stand"],
    ]
    table_slide(s, rows, Inches(0.5), Inches(1.35), Inches(12.3), 0.48, [Inches(2.4), Inches(4.95), Inches(4.95)], font_size=15)
    bullets(
        s,
        Inches(0.5),
        Inches(4.15),
        Inches(12.3),
        Inches(2.8),
        [
            "已排除 kick / uppercut / wing_chun 等攻击性 ActionGroup。",
            "调度状态机不变：投票稳定 → 播动作 → stand 恢复 → 冷却 → idle。",
            "Mac facedetect_mac_demo 只模拟时长，不驱动舵机；真机仍走 FaceDetect。",
            "改任何 B 侧代码前必须先跑 A 回归：失败则停止优化、先修回基线。",
        ],
        16,
    )

    # 6 cloud experiment setup
    s = new_slide(prs)
    header(s, "云侧联调实验：测的是什么", "入口 pipeline.run_mac_e2e_timed　·　机器 Mac M4　·　2026-08-18")
    bullets(
        s,
        Inches(0.5),
        Inches(1.3),
        Inches(12.3),
        Inches(2.2),
        [
            "目的：把「摄像头 + 麦克风 → 云 LLM → MoMask prompt/关节」在 Mac 上跑通并记 CSV。",
            "Decide：gpt-4o-mini @ dmxapi；ASR：最终用 gpt-4o-transcribe（whisper-1 网关 404）。",
            "视觉模型仍在本机：FaceMesh + 几何逻辑回归，帧不上传云。",
            "三组实验：① 实机 cam+mic dry-run　② 实机 cam+mic 全链路　③ mock 感知 + 真 Decide + 真 MoMask。",
        ],
        16,
    )
    rows = [
        ["实验", "感知", "Decide", "MoMask", "总计", "结果"],
        ["实机 + ASR 404", "10.58s", "2.19s", "未跑", "15.0s", "降级 A_only"],
        ["实机 + 可用 ASR", "9.95s", "2.61s", "未跑", "17.4s", "conf=0.20 fallback"],
        ["mock + 云 Decide + MoMask", "~0", "2.75s", "8.88s", "11.6s", "出 joints，未降级"],
    ]
    table_slide(
        s,
        rows,
        Inches(0.4),
        Inches(3.7),
        Inches(12.5),
        0.42,
        [Inches(3.3), Inches(1.7), Inches(1.6), Inches(1.6), Inches(1.5), Inches(2.8)],
        font_size=13,
    )

    # 7 why latency
    s = new_slide(prs)
    header(s, "云侧总时延高：最大头不是 LLM", "实机 15–17s 里 MoMask 甚至没跑；墙钟被「每轮冷启动摄像头」占满")
    add_text(s, Inches(0.5), Inches(1.25), Inches(12.3), Inches(0.35), "实机一轮墙钟构成（感知并行，其余串行）", 15, True, NAVY)
    bars = [
        ("视觉（本机标定+采样）", 10.3, NAVY, "约 10s　·　最大头"),
        ("录音 3.5s（与视觉并行，被盖住）", 3.5, MUTED, "不增加墙钟"),
        ("云 ASR", 3.5, ACCENT, "约 2.2–4.8s"),
        ("云 Decide", 2.5, GOLD, "约 2.2–2.6s"),
        ("MoMask（实机两轮）", 0.0, RED, "降级未跑"),
    ]
    maxv = 10.3
    for i, (name, v, col, note) in enumerate(bars):
        y = Inches(1.7 + i * 0.72)
        add_text(s, Inches(0.5), y, Inches(4.3), Inches(0.4), name, 14, False, INK)
        bw = Inches(0.15 + 6.2 * (v / maxv) if v > 0 else 0.15)
        add_rect(s, Inches(4.9), y + Inches(0.08), bw, Inches(0.38), col)
        add_text(s, Inches(4.95) + bw, y, Inches(3.5), Inches(0.45), note, 13, False, MUTED)
    add_text(
        s,
        Inches(0.5),
        Inches(5.45),
        Inches(12.3),
        Inches(1.5),
        "结论：真正的云开销只有 ASR + Decide（合计约 5–7s）。\n"
        "15s 级体验差，首先应改「每轮重新开摄像头并等待 8s 标定」，而不是先换更大的云模型。",
        16,
        False,
        INK,
    )

    # 8 vision frames
    s = new_slide(prs)
    header(s, "视觉细节：本机几何模型 + 帧数", "capture_perception_mac　·　帧不进云　·　云 Decide 只吃 JSON 标签")
    rows = [
        ["阶段", "帧数 / 时间", "做什么"],
        ["FaceMesh warmup", "15 帧", "跳过开相机后曝光/对焦未稳的帧"],
        ["中性基线标定", "calibration_frames = 20", "采放松脸；相对基线算几何偏移"],
        ["外层等待上限", "最多 8.0 s", "标定未完成也强制结束 → 视觉墙钟≈10s 的主因"],
        ["正式采样", "sample_frames = 12", "多数投票出情绪；同标签分数取平均"],
        ["帧间隔", "sleep 30 ms", "避免连读同一帧"],
    ]
    table_slide(
        s,
        rows,
        Inches(0.4),
        Inches(1.28),
        Inches(12.5),
        0.42,
        [Inches(2.8), Inches(3.6), Inches(6.1)],
        font_size=14,
    )
    bullets(
        s,
        Inches(0.5),
        Inches(4.05),
        Inches(12.3),
        Inches(2.9),
        [
            "模型：MediaPipe FaceMesh 468 点 → 几何特征 → 多项式逻辑回归 4 类（neutral/happy/unhappy/surprised）。",
            "合计帧量级：15 + 20 + 12 ≈ 47 帧；10s 不是「12 帧推理很慢」，而是「每轮冷启动 + 8s 标定超时」。",
            "置信度 ≥ 0.55 → strong，否则 mild。标定不稳时常落到 neutral/0.20 → Arbiter 降级只 A。",
            "对比：实时窗口脚本摄像头常开、逐帧刷新；E2E 计时脚本每轮重新打开相机，两者不可比。",
        ],
        15,
    )

    # 9 audio
    s = new_slide(prs)
    header(s, "听觉细节：录音并行、ASR 串行、不做语音情绪", "云侧只把转写当意图，不把波形送 MoMask")
    stages = [
        ("1 并行录音", "与视觉同时开始\n默认 4s，实测 --audio-sec 3.5\n本机 PyAudio / ffmpeg"),
        ("2 云端 ASR", "视觉结束后串行调用\ngpt-4o-transcribe\n实测 2.2–4.8s"),
        ("3 写入 Perception", "transcript 进 JSON\n无独立 SER 模型\n意图跟语音走"),
        ("4 交给 Decide", "与面部情绪拼 prompt\n有明确需求则跟语音\n脸部只定语气强弱"),
    ]
    for i, (t, d) in enumerate(stages):
        x = Inches(0.45 + i * 3.2)
        add_rect(s, x, Inches(1.4), Inches(3.0), Inches(2.7), LIGHT)
        add_rect(s, x, Inches(1.4), Inches(3.0), Inches(0.55), ACCENT)
        add_text(s, x, Inches(1.5), Inches(3.0), Inches(0.4), t, 16, True, WHITE, PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.12), Inches(2.15), Inches(2.75), Inches(1.8), d, 14, False, INK, PP_ALIGN.CENTER)
    bullets(
        s,
        Inches(0.5),
        Inches(4.35),
        Inches(12.3),
        Inches(2.6),
        [
            "网关坑：whisper-1 → DeploymentNotFound(404)；whisper 无渠道。可用名：gpt-4o-transcribe。",
            "录音 3.5s 被视觉 10s 覆盖，故感知墙钟仍≈10s；ASR 是额外串行成本。",
            "实机一轮转写曾为 “Oh, oh, oh.”，无明确意图 + 弱脸 → fallback、空 prompt、不跑 MoMask。",
            "这是安全规则生效，不是 MoMask 坏了。联调完整 B 应用 mock 感知或明显表情 + 有意图的话。",
        ],
        15,
    )

    # 10 momask 2 vs 8
    s = new_slide(prs)
    header(s, "MoMask：2.3s 并没有变成 8.9s", "同一套权重；差的是「计什么」和「进程是否常驻」")
    rows = [
        ["对比项", "Pi 100 条基准 ≈2.30s", "Mac 云侧 B 路径 8.88s"],
        ["机器", "树莓派 5 CPU 4GB", "Mac M4，--gpu-id -1（CPU）"],
        ["脚本", "bench_pi_gen.py 进程常驻", "track_b 每次 subprocess 新开 Python"],
        ["模型加载", "10.12s，单独记账，不计入条均", "加载 + 生成捆在一次计时里"],
        ["预热", "前 2 条丢掉", "无预热，只跑 1 条"],
        ["计入样本", "98 / 100", "1 条"],
        ["报的数", "纯生成（模型已在内存）", "冷启动解释器 + CLIP + 四权重 + 生成"],
    ]
    table_slide(
        s,
        rows,
        Inches(0.35),
        Inches(1.25),
        Inches(12.6),
        0.42,
        [Inches(2.0), Inches(5.3), Inches(5.3)],
        font_size=13,
    )
    add_rect(s, Inches(0.45), Inches(4.55), Inches(12.4), Inches(2.35), LIGHT)
    add_text(s, Inches(0.7), Inches(4.7), Inches(12), Inches(0.4), "结论（请组会统一口径）", 16, True, RED)
    add_text(
        s,
        Inches(0.7),
        Inches(5.15),
        Inches(12),
        Inches(1.5),
        "2.3s = 热启动生成时延（端侧正式指标）。8.9s = 每轮重新拉起 gen_t2m.py 的进程时延。\n"
        "Pi 自己的加载就是 10.12s，和 Mac 的 8.88s 同属「冷启动」，不能用来否定 2.3s 基线。\n"
        "工程修复：Track B 改为常驻 worker（加载一次、多条生成），Mac/Pi 都应回到 ~2s 档。",
        15,
        False,
        INK,
    )

    # 11 pi bench
    s = new_slide(prs)
    header(s, "端侧 MoMask 基准（4GB Pi，正式生成指标）", "100 条情绪均衡 prompt　·　time_steps=18　·　不计加载")
    kpis = [
        ("2.30s", "平均端到端"),
        ("2.33s", "中位"),
        ("2.64s", "P95"),
        ("1.00", "数值有效率"),
        ("10.12s", "模型加载一次"),
    ]
    for i, (n, lab) in enumerate(kpis):
        x = Inches(0.4 + i * 2.55)
        add_rect(s, x, Inches(1.35), Inches(2.4), Inches(1.7), NAVY if i < 4 else ACCENT)
        add_text(s, x, Inches(1.5), Inches(2.4), Inches(0.7), n, 28, True, WHITE, PP_ALIGN.CENTER)
        add_text(s, x, Inches(2.25), Inches(2.4), Inches(0.5), lab, 13, False, WHITE, PP_ALIGN.CENTER)
    rows = [
        ["项", "数值"],
        ["设备 / 线程", "Raspberry Pi 5 CPU 4GB　/　4"],
        ["总指令 / 预热 / 计入", "100 / 2 / 98"],
        ["最短 / 最长", "1.87s / 2.69s"],
        ["输出", "(T, 22, 3) joints.npy　·　不渲染视频"],
    ]
    table_slide(
        s,
        rows,
        Inches(0.5),
        Inches(3.3),
        Inches(12.3),
        0.42,
        [Inches(4.5), Inches(7.8)],
        font_size=14,
    )

    # 12 emotion metrics
    s = new_slide(prs)
    header(s, "视觉情绪识别：已有离线口径", "几何模型 geometry_emotion_model.json　·　留一人交叉验证")
    kpis = [
        ("0.85", "中性正确率"),
        ("0.73", "happy 召回"),
        ("0.45", "surprised 召回"),
        ("0.46", "unhappy 召回"),
        ("≈0.66", "表情 F1 均值"),
    ]
    for i, (n, lab) in enumerate(kpis):
        x = Inches(0.4 + i * 2.55)
        add_rect(s, x, Inches(1.4), Inches(2.4), Inches(1.55), LIGHT)
        add_rect(s, x, Inches(1.4), Inches(2.4), Inches(0.08), ACCENT)
        add_text(s, x, Inches(1.55), Inches(2.4), Inches(0.7), n, 26, True, NAVY, PP_ALIGN.CENTER)
        add_text(s, x, Inches(2.25), Inches(2.4), Inches(0.5), lab, 13, False, MUTED, PP_ALIGN.CENTER)
    bullets(
        s,
        Inches(0.5),
        Inches(3.25),
        Inches(12.3),
        Inches(3.6),
        [
            "数据：自采约 3265 帧（3 人）+ 公开集；部署模型元数据约 866 人 / 5653 样本。",
            "中性稳、happy 相对最好认；unhappy/surprised 仍弱 —— 这会直接表现为实机 conf 偏低、Arbiter 走 A。",
            "因此 B 的细动作不能只靠 4 类标签：必须靠 Decide 把线索升级成英文动作句。",
            "决策 mock 规则冒烟：e2e_spec 20/20，pass_rate = 1.0（合法 JSON / 字段约束）。",
        ],
        16,
    )

    # 13 decide models
    s = new_slide(prs)
    header(s, "中间 LLM：云侧联调 vs 端侧正式", "端侧不再使用 OpenAI；正式选型沿用此前建议")
    rows = [
        ["环境", "Decide 模型", "用途", "状态"],
        ["Mac 联调", "gpt-4o-mini（中转）", "打通 JSON + prompt 质量", "已跑通，约 2.5s"],
        ["Mac ASR", "gpt-4o-transcribe", "麦克风转写", "已替换不可用的 whisper-1"],
        ["Pi 正式（规划）", "Qwen2.5-1.5B-Instruct Q4 GGUF", "端侧写 action_prompt", "代码已同步，运行时装编译中"],
        ["Pi 4GB 过渡", "规则引擎 edge_rule", "无 GGUF 时的兜底", "A∥B dry-run 已通"],
        ["Pi 16GB 真机", "同 1.5B，可升 3B", "与 MoMask 同机更从容", "待部署"],
    ]
    table_slide(
        s,
        rows,
        Inches(0.35),
        Inches(1.28),
        Inches(12.6),
        0.42,
        [Inches(2.4), Inches(4.0), Inches(3.4), Inches(2.8)],
        font_size=13,
    )
    bullets(
        s,
        Inches(0.5),
        Inches(4.2),
        Inches(12.3),
        Inches(2.7),
        [
            "引擎：llama.cpp / llama-cpp-python；权重国内走 ModelScope / HF 镜像。",
            "4GB 测试机：LLM 与 MoMask 必须错峰；16GB 才是「同机正式默认」。",
            "不要默认 7B：16GB 能塞下，但常见只有 2–3 tok/s，拖垮感知→动作节奏。",
        ],
        16,
    )

    # 14 pi deploy memory
    s = new_slide(prs)
    header(s, "4GB 树莓派部署：进度与内存课", "SSH 已通　·　规则 A∥B 已通　·　Qwen 运行时需本机编译")
    bullets(
        s,
        Inches(0.5),
        Inches(1.3),
        Inches(12.3),
        Inches(2.4),
        [
            "已完成：pipeline / decide_edge 同步；A_parallel_B + 规则 Decide dry-run（A=wave，B 写出 prompt）。",
            "健康：重启后 throttled=0x0，空闲 available ≈ 3GiB，温度空闲 ~32°C。",
            "阻塞：Python 3.13 + aarch64 无现成 llama-cpp-python wheel → 本机 g++ 编译。",
        ],
        16,
    )
    add_rect(s, Inches(0.45), Inches(3.7), Inches(12.4), Inches(3.15), LIGHT)
    add_text(s, Inches(0.7), Inches(3.85), Inches(12), Inches(0.4), "为什么「安装」也会 OOM？（不是只有推理才占 RAM）", 16, True, RED)
    add_text(
        s,
        Inches(0.7),
        Inches(4.35),
        Inches(12),
        Inches(2.2),
        "4GB 是运行内存。编译时多个 cc1plus（-O3）每个可占数百 MB；默认多核并行峰值 2–3GB+，再叠加桌面即顶满。\n"
        "实测：可用内存一度降至 ~440MB 并开始用 swap，温度 ~68°C。已改为 CMAKE_BUILD_PARALLEL_LEVEL=1。\n"
        "16GB 同规格机器：编译 OOM 基本不会再发生；编译仍会慢，但机器不容易被编死。\n"
        "推理阶段：Qwen Q4 约 1–2GB，MoMask 约 2GB+，4GB 上两者仍建议错峰，16GB 可同机。",
        14,
        False,
        INK,
    )

    # 15 next
    s = new_slide(prs)
    header(s, "下一步（按优先级）", "先统一延迟口径，再补端侧 LLM，再交关节包")
    steps = [
        ("P0", "常驻 MoMask worker", "subprocess 冷启动改为加载一次；让 Mac/Pi 生成时延回到 ~2s 口径。"),
        ("P0", "摄像头常开标定一次", "E2E 不再每轮 8s 标定；视觉应从 10s 降到百毫秒级采样。"),
        ("P1", "装完 Qwen 1.5B 并冒烟", "单线程编完 llama-cpp → ModelScope 拉 GGUF → edge_llm Decide。"),
        ("P1", "收紧 prompt 句式", "强制 a person … HumanML3D；弱脸但有 ASR 意图时仍出 prompt。"),
        ("P2", "端侧 ASR", "云 Whisper 仅 Mac 联调；Pi 上换 whisper.cpp tiny/base 或先手写 transcript。"),
        ("P2", "16GB 真机 + 队友关节", "同机 A∥B；joints 包合同与舵机映射联调。"),
    ]
    for i, (tag, t, d) in enumerate(steps):
        y = Inches(1.28 + i * 0.88)
        add_rect(s, Inches(0.45), y, Inches(12.4), Inches(0.78), LIGHT)
        add_rect(s, Inches(0.45), y, Inches(1.15), Inches(0.78), NAVY if tag == "P0" else ACCENT)
        add_text(s, Inches(0.45), y + Inches(0.2), Inches(1.15), Inches(0.4), tag, 16, True, WHITE, PP_ALIGN.CENTER)
        add_text(s, Inches(1.8), y + Inches(0.08), Inches(10.8), Inches(0.32), t, 16, True, NAVY)
        add_text(s, Inches(1.8), y + Inches(0.4), Inches(10.8), Inches(0.32), d, 13, False, MUTED)

    # 16 takeaway
    s = new_slide(prs)
    header(s, "一页结论（可直接贴纪要）", "")
    takes = [
        "架构已冻结：A 保底、B 正式、默认 A∥B；降级只 A、不撤销 A。",
        "A 轨 enrichment 离线全过；B 轨 Mac 联调已通，决策 mock 20/20。",
        "云侧 15s 级墙钟：~10s 是每轮摄像头标定，云只贡献 ASR+Decide≈5–7s。",
        "视觉：本机 FaceMesh，约 15+20+12 帧；12 帧投票，8s 为标定上限。帧不上云。",
        "听觉：3.5s 录音与视觉并行，随后云 ASR；无语音情绪模型。",
        "MoMask 正式指标仍是 Pi 热启动 2.30s/条；8.9s 是 Mac 冷启动进程，不是模型变慢。",
        "端侧 Decide 正式用 Qwen2.5-1.5B；4GB 需单线程编译且与 MoMask 错峰，16GB 才同机。",
    ]
    for i, t in enumerate(takes):
        y = Inches(1.28 + i * 0.72)
        n = "%d" % (i + 1)
        add_rect(s, Inches(0.5), y, Inches(0.55), Inches(0.55), GOLD if i in (2, 5) else ACCENT)
        add_text(s, Inches(0.5), y + Inches(0.08), Inches(0.55), Inches(0.4), n, 16, True, WHITE, PP_ALIGN.CENTER)
        add_text(s, Inches(1.2), y + Inches(0.08), Inches(11.5), Inches(0.5), t, 16, False, INK)

    # end
    s = new_slide(prs)
    add_rect(s, 0, 0, W, H, NAVY)
    add_rect(s, Inches(0.55), Inches(2.7), Inches(1.4), Inches(0.08), GOLD)
    add_text(s, Inches(0.55), Inches(2.95), Inches(12), Inches(0.8), "谢谢。欢迎讨论口径与下一步排期。", 32, True, WHITE)
    add_text(
        s,
        Inches(0.55),
        Inches(4.0),
        Inches(12),
        Inches(1.2),
        "建议先对齐三句话：视觉 10s 是标定；MoMask 2.3s 是热启动；8.9s 是冷启动。",
        18,
        False,
        RGBColor(0xC5, 0xD4, 0xE0),
    )
    add_text(s, Inches(0.55), Inches(6.3), Inches(12), Inches(0.4), "材料：pipeline_runs 计时 CSV  ·  experiment/pi_bench  ·  BASELINE_FREEZE.md", 14, False, RGBColor(0xA8, 0xBC, 0xCC))

    prs.save(str(OUT_PPTX))
    print("wrote", OUT_PPTX)


if __name__ == "__main__":
    build()
