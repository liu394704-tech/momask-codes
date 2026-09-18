#!/usr/bin/env python3
"""组会 PPT：两页。整体流程延迟 + 当前问题。"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path("/Users/emmaliu/Desktop/HKUST/RBM-project/Model/momask-codes")
OUT = ROOT / "experiment/e2e_spec/组会汇报_决策延迟复测_20260909.pptx"

INK = RGBColor(0x1B, 0x1B, 0x1B)
MUTED = RGBColor(0x5C, 0x5C, 0x5C)
NAVY = RGBColor(0x1B, 0x3A, 0x4B)
TEAL = RGBColor(0x2A, 0x6F, 0x6F)
CORAL = RGBColor(0xC4, 0x5C, 0x41)
CREAM = RGBColor(0xF7, 0xF4, 0xEE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE6, 0xE0, 0xD6)
FONT = "PingFang SC"
W, H = Inches(13.333), Inches(7.5)


def font(run, text, size, bold=False, color=INK):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", FONT)


def rect(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def txt(slide, l, t, w, h, text, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    font(r, text, size, bold, color)
    return box


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, CREAM)
    return s


def bar(slide):
    rect(slide, 0, 0, Inches(0.12), H, NAVY)


def foot(slide, page):
    txt(slide, Inches(0.55), Inches(7.12), Inches(10.4), Inches(0.28),
        "组会  ·  感知 → 决策 → 动作  ·  2026.09.09 Pi 复测", 11, False, MUTED)
    txt(slide, Inches(11.6), Inches(7.12), Inches(1.2), Inches(0.28),
        "%d / 2" % page, 11, False, MUTED, PP_ALIGN.RIGHT)


def title_block(slide, kicker, title, sub=""):
    bar(slide)
    txt(slide, Inches(0.5), Inches(0.22), Inches(12), Inches(0.26), kicker, 12, False, TEAL)
    txt(slide, Inches(0.5), Inches(0.46), Inches(12.3), Inches(0.42), title, 24, True, NAVY)
    if sub:
        txt(slide, Inches(0.5), Inches(0.9), Inches(12.3), Inches(0.32), sub, 13, False, MUTED)
    rect(slide, Inches(0.5), Inches(1.24), Inches(1.1), Inches(0.05), CORAL)


def card(slide, l, t, w, h, fill=WHITE):
    sh = rect(slide, l, t, w, h, fill)
    sh.line.fill.solid()
    sh.line.color.rgb = LINE
    sh.line.width = Pt(1)
    return sh


def table(slide, l, t, w, h, rows, col_w=None):
    n_row, n_col = len(rows), len(rows[0])
    sh = slide.shapes.add_table(n_row, n_col, l, t, w, h)
    tbl = sh.table
    if col_w:
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = cw
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = p.add_run()
            font(r, val, 12 if i else 11, bold=(i == 0 or j == 0),
                 color=WHITE if i == 0 else INK)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if i == 0 else (CREAM if i % 2 else WHITE)
    return sh


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # -------- page 1: pipeline + latency --------
    s = blank(prs)
    title_block(
        s, "整体流程",
        "从前到后每一步的延迟",
        "2026.09.09 09:58 在 Raspberry Pi 5（10.160.13.84）复测决策；其余步骤标注数据来源。",
    )
    foot(s, 1)

    table(
        s, Inches(0.5), Inches(1.42), Inches(12.35), Inches(4.05),
        [
            ["步骤", "做什么", "延迟", "数据从哪来"],
            ["① 摄像头", "人脸表情 → happy / frown 等文字", "约 0.1–0.3s", "板子本地算法；今日未复测"],
            ["② 麦克风", "语音 → 命令字 / 转写文字", "硬件词条约几十 ms", "真机 WonderEcho；今日未复测"],
            ["③ 决策", "文字融合 → 一句 action_prompt", "云端 0.80s（最慢 1.33s）", "今日 Pi 实测 6 次"],
            ["", "同上，改走本地开源模型", "1.5B：1.51s　0.5B：0.59s", "今日 Pi 实测，断网也能跑"],
            ["④ 先动一下", "挥手 / 鞠躬等预设动作立刻执行", "< 0.2s 开始动", "工程已具备，遮住等待"],
            ["⑤ MoMask", "prompt → 关节 (T,22,3)", "Pi 上约 2s（估计）", "今日未在 Pi 复测"],
        ],
        col_w=[Inches(1.7), Inches(3.9), Inches(3.55), Inches(3.2)],
    )

    card(s, Inches(0.5), Inches(5.58), Inches(12.35), Inches(1.4))
    txt(s, Inches(0.72), Inches(5.7), Inches(12), Inches(1.15),
        "合计（串行）：决策 0.80 + MoMask≈2 ≈ 3 秒量级，接近 3–4 秒目标。\n"
        "对人的体感：④ 先动，1 秒内机器人已经有反应；完整生成动作稍后补上。\n"
        "今日确认的只有第③步。①②⑤ 仍是已有结论或估计，不能当成今天刚测死的数。",
        14, False, INK)

    # -------- page 2: problems --------
    s = blank(prs)
    title_block(
        s, "当前问题",
        "现在卡在哪",
        "决策速度已经有数；真正还没闭环的是质量和整机。",
    )
    foot(s, 2)

    problems = [
        ("1 决策绑在网上会抖",
         "Flash 中位 0.80s，但最慢到过 1.33s；断网这次请求直接失败。不能把“机器人能不能动”绑死在云端。"),
        ("2 本地模型快了，质量一般",
         "1.5B 能出合法 JSON，但情绪偶发判错、动作描述会越界（例如让机器人坐下）。0.5B 更快，语义更糙。"),
        ("3 MoMask 在 Pi 上还没今日实测",
         "完整动作时间仍按约 2s 估计。若实际是 4–5s，瓶颈就不在决策，而在生成。这是下一个必须补的数。"),
        ("4 真机全链路还没串起来",
         "表情、硬件麦克风、决策、预设动作、MoMask 目前是分段验证。还没有一次「真人面对机器人、从看到动」的计时。"),
        ("5 闭源 / 开源容易讲混",
         "Flash 快，但是云端；Qwen2.5 才是板子上跑。Pi 联网可以用闭源，但不等于模型装在端侧。"),
    ]
    y = 1.42
    for title, body in problems:
        card(s, Inches(0.5), Inches(y), Inches(12.35), Inches(1.0))
        txt(s, Inches(0.7), Inches(y + 0.08), Inches(12), Inches(0.32), title, 15, True, NAVY)
        txt(s, Inches(0.7), Inches(y + 0.42), Inches(12), Inches(0.5), body, 13, False, INK)
        y += 1.08

    prs.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
